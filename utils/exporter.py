from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EXPORTS_DIR = Path(__file__).parent.parent / "exports"

# ── Palette ──────────────────────────────────────────────────────────────────
BG_TITLE      = "1F3864"   # dark navy
BG_SECTION    = "2E75B6"   # medium blue
BG_META_LABEL = "D6DCE4"   # light slate
BG_PASS       = "E2EFDA"   # soft green
BG_MINOR      = "FFF2CC"   # soft amber
BG_MAJOR      = "FFE0E0"   # soft red
BG_ALT        = "F5F5F5"   # alternating row grey
BG_TBL_HDR    = "BDD7EE"   # table header blue

FG_WHITE  = "FFFFFF"
FG_PASS   = "375623"
FG_MINOR  = "7F6000"
FG_MAJOR  = "8B0000"
FG_DARK   = "1A1A1A"

BORDER_COLOR = "BFBFBF"

STATUS_STYLE = {
    "Pass":              (BG_PASS,  FG_PASS),
    "Needs Minor Edits": (BG_MINOR, FG_MINOR),
    "Needs Major Edits": (BG_MAJOR, FG_MAJOR),
    "Error":             (BG_MAJOR, FG_MAJOR),
}

SEVERITY_STYLE = {
    "major": (BG_MAJOR, FG_MAJOR),
    "minor": (BG_MINOR, FG_MINOR),
}

LAST_COL = "E"
NUM_COLS = 5   # A-E


# ── Style helpers ─────────────────────────────────────────────────────────────
def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _font(bold=False, color=FG_DARK, size=11, italic=False) -> Font:
    return Font(name="Calibri", bold=bold, color=color, size=size, italic=italic)


def _align(wrap=False, h="left", v="center") -> Alignment:
    return Alignment(wrap_text=wrap, horizontal=h, vertical=v)


def _border_side(style="thin") -> Side:
    return Side(style=style, color=BORDER_COLOR)


def _thin_border() -> Border:
    s = _border_side()
    return Border(left=s, right=s, top=s, bottom=s)


# ── Main Exporter ─────────────────────────────────────────────────────────────
class Exporter:
    def to_csv(self, result: dict, client_id: str, stage: str,
               client_name: str = None) -> str:
        EXPORTS_DIR.mkdir(exist_ok=True)
        ts = datetime.now()
        ts_file    = ts.strftime("%Y%m%d_%H%M%S")
        ts_display = ts.strftime("%B %d, %Y  %I:%M %p")

        filename = f"{client_id}_{stage}_{ts_file}.xlsx"
        filepath = EXPORTS_DIR / filename

        wb = Workbook()
        ws = wb.active
        ws.title = "QC Report"

        # Column widths
        widths = [22, 14, 30, 52, 40]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        r = 1  # current row pointer

        # ── Report Title ──────────────────────────────────────────────────────
        r = self._title_row(ws, r, "GBP Quality Check Report")

        # ── Metadata block ────────────────────────────────────────────────────
        stage_label = stage.replace("_", " ").title()
        score       = result.get("qc_score", 0)
        status      = result.get("overall_status", "")
        status_bg, status_fg = STATUS_STYLE.get(status, ("FFFFFF", FG_DARK))

        r = self._meta_row(ws, r, "Client",  client_name or client_id,
                                  "Stage",   stage_label)
        r = self._meta_row(ws, r, "Date",    ts_display,
                                  "QC Score", f"{score} / 100")

        # Status spans full width
        ws[f"A{r}"] = "Status"
        ws[f"A{r}"].font  = _font(bold=True)
        ws[f"A{r}"].fill  = _fill(BG_META_LABEL)
        ws[f"A{r}"].alignment = _align()
        ws.merge_cells(f"B{r}:{LAST_COL}{r}")
        ws[f"B{r}"] = status
        ws[f"B{r}"].font  = _font(bold=True, color=status_fg, size=13)
        ws[f"B{r}"].fill  = _fill(status_bg)
        ws[f"B{r}"].alignment = _align(h="center")
        ws.row_dimensions[r].height = 24
        r += 1

        r = self._spacer(ws, r)

        # ── Passed Checks ─────────────────────────────────────────────────────
        r = self._section_header(ws, r, "PASSED CHECKS")
        passed = result.get("passed_checks", [])
        if passed:
            r = self._col_header(ws, r, ["Check"], span_to=LAST_COL)
            for i, check in enumerate(passed):
                bg = "FFFFFF" if i % 2 == 0 else BG_ALT
                ws.merge_cells(f"A{r}:{LAST_COL}{r}")
                c = ws[f"A{r}"]
                c.value     = f"OK   {check}"
                c.font      = _font(color=FG_PASS)
                c.fill      = _fill(bg)
                c.border    = _thin_border()
                c.alignment = _align()
                r += 1
        else:
            r = self._empty_row(ws, r, "No checks passed.")

        r = self._spacer(ws, r)

        # ── Issues Found ──────────────────────────────────────────────────────
        r = self._section_header(ws, r, "ISSUES FOUND")
        issues = result.get("issues_found", [])
        if issues:
            r = self._col_header(ws, r,
                                 ["Severity", "Check", "Description",
                                  "Guideline / Rule Violated"])
            for issue in issues:
                sev = issue.get("severity", "minor")
                ib, ifg = SEVERITY_STYLE.get(sev, ("FFFFFF", FG_DARK))

                def _icell(col, val, bold=False, wrap=False, center=False):
                    c = ws[f"{col}{r}"]
                    c.value     = val
                    c.font      = _font(bold=bold, color=ifg, italic=not bold)
                    c.fill      = _fill(ib)
                    c.border    = _thin_border()
                    c.alignment = _align(wrap=wrap,
                                         h="center" if center else "left")

                _icell("A", sev.upper(), bold=True, center=True)
                _icell("B", issue.get("check", ""), bold=True)
                _icell("C", issue.get("issue", ""), wrap=True)
                _icell("D", issue.get("reason", ""), wrap=True)
                ws.row_dimensions[r].height = 48
                r += 1
        else:
            r = self._empty_row(ws, r, "No issues found.", green=True)

        r = self._spacer(ws, r)

        # ── Recommended Fixes ─────────────────────────────────────────────────
        r = self._section_header(ws, r, "RECOMMENDED FIXES")
        fixes = result.get("recommended_fixes", [])
        if fixes:
            r = self._col_header(ws, r, ["#", "Fix"], col_b_span=True)
            for i, fix in enumerate(fixes, 1):
                ws[f"A{r}"] = str(i)
                ws[f"A{r}"].font      = _font(bold=True)
                ws[f"A{r}"].fill      = _fill(BG_META_LABEL)
                ws[f"A{r}"].border    = _thin_border()
                ws[f"A{r}"].alignment = _align(h="center")
                ws.merge_cells(f"B{r}:{LAST_COL}{r}")
                ws[f"B{r}"] = fix
                ws[f"B{r}"].font      = _font()
                ws[f"B{r}"].border    = _thin_border()
                ws[f"B{r}"].alignment = _align(wrap=True)
                ws.row_dimensions[r].height = 36
                r += 1
        else:
            r = self._empty_row(ws, r, "No fixes required.", green=True)

        r = self._spacer(ws, r)

        # ── Recommended Content (Suggested Edited Version) ────────────────────
        r = self._section_header(ws, r, "RECOMMENDED CONTENT — Suggested Edited Version")
        edited = (result.get("suggested_edited_version") or "").strip()
        if edited:
            char_count = len(edited)
            ws.merge_cells(f"A{r}:{LAST_COL}{r}")
            ws[f"A{r}"] = edited
            ws[f"A{r}"].font      = _font(size=11)
            ws[f"A{r}"].border    = _thin_border()
            ws[f"A{r}"].alignment = _align(wrap=True, v="top")
            # Height: roughly 15px per 80 chars, min 80, max 400
            ws.row_dimensions[r].height = max(80, min(char_count // 5, 400))
            r += 1
            ws.merge_cells(f"A{r}:{LAST_COL}{r}")
            ws[f"A{r}"] = f"{char_count} characters"
            ws[f"A{r}"].font      = _font(italic=True, color="888888")
            ws[f"A{r}"].fill      = _fill(BG_ALT)
            ws[f"A{r}"].alignment = _align(h="right")
            r += 1
        else:
            r = self._empty_row(ws, r,
                                "No edits required — original content approved.",
                                green=True)

        r = self._spacer(ws, r)

        # ── Final Approval Recommendation ─────────────────────────────────────
        r = self._section_header(ws, r, "FINAL APPROVAL RECOMMENDATION")
        rec = (result.get("final_approval_recommendation") or "").strip()
        if rec:
            ws.merge_cells(f"A{r}:{LAST_COL}{r}")
            status_bg2, status_fg2 = STATUS_STYLE.get(status, ("FFFFFF", FG_DARK))
            ws[f"A{r}"] = rec
            ws[f"A{r}"].font      = _font(color=status_fg2)
            ws[f"A{r}"].fill      = _fill(status_bg2)
            ws[f"A{r}"].border    = _thin_border()
            ws[f"A{r}"].alignment = _align(wrap=True, v="top")
            ws.row_dimensions[r].height = 60

        wb.save(str(filepath))
        return str(filepath)

    def to_generated_excel(self, result: dict, client_id: str, location: str,
                           client_name: str = None) -> str:
        EXPORTS_DIR.mkdir(exist_ok=True)
        ts = datetime.now()
        ts_file    = ts.strftime("%Y%m%d_%H%M%S")
        ts_display = ts.strftime("%B %d, %Y  %I:%M %p")
        safe_loc   = location.replace(" ", "_").replace(",", "")
        filename   = f"{client_id}_generated_{safe_loc}_{ts_file}.xlsx"
        filepath   = EXPORTS_DIR / filename

        wb = Workbook()
        ws = wb.active
        ws.title = "Generated Content"

        widths = [22, 14, 30, 52, 40]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        r = 1
        r = self._title_row(ws, r, "GBP Location Content — Generated Post")

        char_count = result.get("character_count", 0)
        within     = result.get("within_limit", char_count <= 1500)
        r = self._meta_row(ws, r, "Client",    client_name or client_id,
                                  "Location",  location)
        r = self._meta_row(ws, r, "Date",      ts_display,
                                  "Characters", f"{char_count} / 1500")

        # Within-limit status
        sl_bg = BG_PASS if within else BG_MAJOR
        sl_fg = FG_PASS if within else FG_MAJOR
        ws[f"A{r}"] = "Status"
        ws[f"A{r}"].font  = _font(bold=True)
        ws[f"A{r}"].fill  = _fill(BG_META_LABEL)
        ws[f"A{r}"].alignment = _align()
        ws.merge_cells(f"B{r}:{LAST_COL}{r}")
        ws[f"B{r}"] = "Within character limit" if within else "OVER CHARACTER LIMIT — must be trimmed"
        ws[f"B{r}"].font  = _font(bold=True, color=sl_fg, size=12)
        ws[f"B{r}"].fill  = _fill(sl_bg)
        ws[f"B{r}"].alignment = _align(h="center")
        ws.row_dimensions[r].height = 24
        r += 1

        r = self._spacer(ws, r)

        # Sections included
        r = self._section_header(ws, r, "SECTIONS INCLUDED")
        sections = result.get("sections_included", [])
        ws.merge_cells(f"A{r}:{LAST_COL}{r}")
        ws[f"A{r}"] = "  |  ".join(sections) if sections else "Not specified"
        ws[f"A{r}"].font      = _font(color=FG_PASS)
        ws[f"A{r}"].fill      = _fill(BG_PASS)
        ws[f"A{r}"].alignment = _align(h="center")
        r += 1

        r = self._spacer(ws, r)

        # Generated post
        r = self._section_header(ws, r, "GENERATED POST — Ready to Copy and Use")
        full_post = (result.get("full_post") or "").strip()
        if full_post:
            ws.merge_cells(f"A{r}:{LAST_COL}{r}")
            ws[f"A{r}"] = full_post
            ws[f"A{r}"].font      = _font(size=11)
            ws[f"A{r}"].border    = _thin_border()
            ws[f"A{r}"].alignment = _align(wrap=True, v="top")
            ws.row_dimensions[r].height = max(120, min(char_count // 4, 500))
            r += 1
            ws.merge_cells(f"A{r}:{LAST_COL}{r}")
            ws[f"A{r}"] = f"{char_count} characters"
            ws[f"A{r}"].font      = _font(italic=True, color="888888")
            ws[f"A{r}"].fill      = _fill(BG_ALT)
            ws[f"A{r}"].alignment = _align(h="right")
            r += 1

        r = self._spacer(ws, r)

        # Customization notes
        r = self._section_header(ws, r, "CUSTOMIZATION NOTES")
        notes = result.get("customization_notes", [])
        if notes:
            for i, note in enumerate(notes):
                bg = "FFFFFF" if i % 2 == 0 else BG_ALT
                ws.merge_cells(f"A{r}:{LAST_COL}{r}")
                ws[f"A{r}"] = note
                ws[f"A{r}"].font      = _font()
                ws[f"A{r}"].fill      = _fill(bg)
                ws[f"A{r}"].border    = _thin_border()
                ws[f"A{r}"].alignment = _align(wrap=True)
                r += 1
        else:
            r = self._empty_row(ws, r, "No notes provided.")

        wb.save(str(filepath))
        return str(filepath)

    def to_all_locations_excel(self, results: dict, client_id: str,
                               post_type: str, client_name: str = None) -> str:
        """Export all generated location posts into a single sheet."""
        EXPORTS_DIR.mkdir(exist_ok=True)
        ts = datetime.now()
        ts_file    = ts.strftime("%Y%m%d_%H%M%S")
        ts_display = ts.strftime("%B %d, %Y  %I:%M %p")
        filename   = f"{client_id}_all_locations_{ts_file}.xlsx"
        filepath   = EXPORTS_DIR / filename

        wb = Workbook()
        ws = wb.active
        ws.title = "All Locations"

        # Two columns: narrow label col + wide post col
        ws.column_dimensions["A"].width = 32   # Location / label
        ws.column_dimensions["B"].width = 100  # Post content

        locations = list(results.keys())
        total     = len(locations)
        r = 1

        # ── Report header ─────────────────────────────────────────────────────
        ws.merge_cells(f"A{r}:B{r}")
        ws[f"A{r}"] = "GBP Location Content — All Locations"
        ws[f"A{r}"].font      = _font(bold=True, color=FG_WHITE, size=16)
        ws[f"A{r}"].fill      = _fill(BG_TITLE)
        ws[f"A{r}"].alignment = _align(h="center")
        ws.row_dimensions[r].height = 34
        r += 1

        for label, val in [
            ("Client",           client_name or client_id),
            ("Post Type",        post_type),
            ("Date",             ts_display),
            ("Total Locations",  str(total)),
        ]:
            ws[f"A{r}"] = label
            ws[f"A{r}"].font  = _font(bold=True)
            ws[f"A{r}"].fill  = _fill(BG_META_LABEL)
            ws[f"A{r}"].alignment = _align()
            ws[f"B{r}"] = val
            ws[f"B{r}"].font  = _font()
            ws[f"B{r}"].alignment = _align()
            ws.row_dimensions[r].height = 18
            r += 1

        r += 1  # spacer

        # ── Column headers ────────────────────────────────────────────────────
        for col, hdr in [("A", "Location"), ("B", "Generated Post")]:
            ws[f"{col}{r}"] = hdr
            ws[f"{col}{r}"].font      = _font(bold=True)
            ws[f"{col}{r}"].fill      = _fill(BG_TBL_HDR)
            ws[f"{col}{r}"].border    = _thin_border()
            ws[f"{col}{r}"].alignment = _align(h="center")
        ws.row_dimensions[r].height = 22
        r += 1

        # ── One row per location ──────────────────────────────────────────────
        for i, loc in enumerate(locations):
            res       = results[loc]
            chars     = res.get("character_count", 0)
            within    = res.get("within_limit", chars <= 1500)
            full_post = (res.get("full_post") or "").strip()
            notes     = res.get("customization_notes", [])

            # Alternate light background for readability
            row_bg = "FFFFFF" if i % 2 == 0 else "EEF4FB"

            # Location cell (col A): name + char count + status
            sl_fg   = FG_PASS if within else FG_MAJOR
            status  = f"{chars} chars  |  {'OK' if within else 'OVER LIMIT'}"
            loc_text = f"{loc}\n{status}"

            ws[f"A{r}"] = loc_text
            ws[f"A{r}"].font      = _font(bold=True, color=sl_fg)
            ws[f"A{r}"].fill      = _fill(BG_PASS if within else BG_MAJOR)
            ws[f"A{r}"].border    = _thin_border()
            ws[f"A{r}"].alignment = _align(wrap=True, v="top")

            # Post content cell (col B)
            ws[f"B{r}"] = full_post
            ws[f"B{r}"].font      = _font(size=11)
            ws[f"B{r}"].fill      = _fill(row_bg)
            ws[f"B{r}"].border    = _thin_border()
            ws[f"B{r}"].alignment = _align(wrap=True, v="top")

            # Row height: enough to show the full post
            ws.row_dimensions[r].height = max(150, min(chars // 3, 400))
            r += 1

        wb.save(str(filepath))
        return str(filepath)

    # ── Layout helpers ────────────────────────────────────────────────────────
    def _title_row(self, ws, r: int, text: str) -> int:
        ws.merge_cells(f"A{r}:{LAST_COL}{r}")
        c = ws[f"A{r}"]
        c.value     = text
        c.font      = _font(bold=True, color=FG_WHITE, size=16)
        c.fill      = _fill(BG_TITLE)
        c.alignment = _align(h="center")
        ws.row_dimensions[r].height = 34
        return r + 1

    def _section_header(self, ws, r: int, title: str) -> int:
        ws.merge_cells(f"A{r}:{LAST_COL}{r}")
        c = ws[f"A{r}"]
        c.value     = title
        c.font      = _font(bold=True, color=FG_WHITE, size=12)
        c.fill      = _fill(BG_SECTION)
        c.alignment = _align(h="left")
        ws.row_dimensions[r].height = 22
        return r + 1

    def _meta_row(self, ws, r: int,
                  label1: str, val1: str,
                  label2: str, val2: str) -> int:
        for col, text, bold, bg in [
            ("A", label1, True,  BG_META_LABEL),
            ("B", val1,   False, "FFFFFF"),
            ("C", label2, True,  BG_META_LABEL),
            ("D", val2,   True,  "FFFFFF"),
        ]:
            ws[f"{col}{r}"] = text
            ws[f"{col}{r}"].font      = _font(bold=bold)
            ws[f"{col}{r}"].fill      = _fill(bg)
            ws[f"{col}{r}"].alignment = _align()
        ws.row_dimensions[r].height = 18
        return r + 1

    def _col_header(self, ws, r: int, labels: list,
                    span_to: str = None, col_b_span: bool = False) -> int:
        cols = "ABCDE"
        for i, label in enumerate(labels):
            col = cols[i]
            if col_b_span and col == "B":
                ws.merge_cells(f"B{r}:{LAST_COL}{r}")
            ws[f"{col}{r}"] = label
            ws[f"{col}{r}"].font      = _font(bold=True)
            ws[f"{col}{r}"].fill      = _fill(BG_TBL_HDR)
            ws[f"{col}{r}"].border    = _thin_border()
            ws[f"{col}{r}"].alignment = _align(h="center")
        if span_to and len(labels) == 1:
            ws.merge_cells(f"A{r}:{span_to}{r}")
        ws.row_dimensions[r].height = 18
        return r + 1

    def _empty_row(self, ws, r: int, text: str, green: bool = False) -> int:
        ws.merge_cells(f"A{r}:{LAST_COL}{r}")
        ws[f"A{r}"] = text
        ws[f"A{r}"].font      = _font(italic=True, color=FG_PASS if green else FG_DARK)
        ws[f"A{r}"].fill      = _fill(BG_PASS if green else BG_ALT)
        ws[f"A{r}"].alignment = _align()
        return r + 1

    def _spacer(self, ws, r: int) -> int:
        ws.row_dimensions[r].height = 10
        return r + 1
